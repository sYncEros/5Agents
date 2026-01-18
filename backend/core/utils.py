from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Iterable, List


def leer_corpus(archivos: Iterable) -> str:
    textos: List[str] = []
    for archivo in archivos:
        nombre = getattr(archivo, "name", "")
        contenido = archivo.read()
        if isinstance(contenido, str):
            textos.append(contenido)
            continue

        if nombre.lower().endswith(".pdf"):
            try:
                from pypdf import PdfReader
            except Exception as exc:  # pragma: no cover
                raise RuntimeError("Instala 'pypdf' para leer PDFs") from exc

            reader = PdfReader(BytesIO(contenido))
            texto_pdf = "\n".join(page.extract_text() or "" for page in reader.pages)
            textos.append(texto_pdf)
        else:
            try:
                textos.append(contenido.decode("utf-8"))
            except Exception:
                textos.append(contenido.decode("latin-1", errors="ignore"))

    return "\n\n".join(t for t in textos if t.strip())


def limpiar_json_output(texto: str) -> str:
    cleaned = texto.strip()
    cleaned = re.sub(r"```json|```", "", cleaned, flags=re.IGNORECASE).strip()
    match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
    if match:
        return match.group(0)
    return cleaned


def validar_paper_markdown(contenido: str) -> bool:
    required_sections = [
        "título",
        "abstract",
        "introducción",
        "desarrollo",
        "discusión",
        "conclusión",
        "referencias",
    ]
    contenido_lower = contenido.lower()
    return all(section in contenido_lower for section in required_sections)


def crear_docx(contenido: str) -> bytes:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("Instala 'python-docx' para exportar a DOCX") from exc

    document = Document()
    for line in contenido.split("\n"):
        if line.startswith("#"):
            level = min(line.count("#"), 4)
            document.add_heading(line.lstrip("# "), level=level)
        else:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()
